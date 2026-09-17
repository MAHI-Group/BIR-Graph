import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

EXCEL_MAX_ROWS = 200_000


def write_excel(path, sheets):
    """sheets: dict of sheet name -> DataFrame. Large tables are truncated
    in the workbook (full versions go to CSV)."""
    header = Font(name="Arial", bold=True, size=10)
    body = Font(name="Arial", size=10)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df = df.head(EXCEL_MAX_ROWS)
            df.to_excel(xw, sheet_name=name[:31], index=False)
            ws = xw.sheets[name[:31]]
            ws.freeze_panes = "A2"
            styled_rows = ws.max_row if ws.max_row * max(ws.max_column, 1) <= 60_000 else 1
            for row in ws.iter_rows(min_row=1, max_row=styled_rows):
                for cell in row:
                    cell.font = header if cell.row == 1 else body
            for k, col in enumerate(df.columns, start=1):
                sample = [str(col)] + [str(v) for v in df[col].head(200)]
                width = min(60, max(8, max(len(s) for s in sample) + 2))
                ws.column_dimensions[get_column_letter(k)].width = width
