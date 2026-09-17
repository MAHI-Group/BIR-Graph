"""Synthetic samples x genes data with planted Boolean implications, used by the tests.

Modules A-H are gene programmes switched by latent Boolean states; N genes are
unstructured noise and L genes have too little dynamic range to be analysed.
"""

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MODULES = {"A": 8, "B": 6, "C": 6, "D": 6, "E": 5, "F": 5, "G": 4, "H": 4}

PLANTED = [
    ("A", "within-module", "genes in a module are Boolean equivalent"),
    ("B high => A high", "HH", "B is switched on only in a subset of A-high samples"),
    ("C opposite D", "OP", "D is high exactly when C is low"),
    ("E high => F low", "HL", "E and F are rarely high together"),
    ("G low => H high", "LH", "G and H are rarely low together"),
    ("A, C, D", "group effect", "A is more often high in Tumour; C in Normal"),
    ("N01-N16", "noise", "unimodal genes with no planted relation"),
    ("L01-L02", "low range", "nearly constant genes; excluded by the dynamic-range filter"),
]


def latent_states(tumour, rng):
    n = tumour.size
    u, v = rng.random(n), rng.random(n)
    a = rng.random(n) < np.where(tumour, 0.7, 0.2)
    c = rng.random(n) < np.where(tumour, 0.35, 0.7)
    return {
        "A": a,
        "B": a & (rng.random(n) < 0.5),
        "C": c,
        "D": ~c,
        "E": u < 0.3,
        "F": (u >= 0.3) & (u < 0.6),
        "G": (v < 0.3) | (v >= 0.6),
        "H": v >= 0.3,
    }


def simulate(n_tumour=150, n_normal=90, flip=0.02, seed=7):
    rng = np.random.default_rng(seed)
    group = np.array(["Tumour"] * n_tumour + ["Normal"] * n_normal)
    rng.shuffle(group)
    n = group.size
    states = latent_states(group == "Tumour", rng)
    data, module_of = {}, {}
    for m, size in MODULES.items():
        for k in range(1, size + 1):
            name = f"{m}{k:02d}"
            on = states[m] ^ (rng.random(n) < flip)
            data[name] = rng.uniform(3.0, 6.0) + rng.uniform(2.5, 4.0) * on + rng.normal(0.0, 0.45, n)
            module_of[name] = m
    for k in range(1, 17):
        data[f"N{k:02d}"] = rng.normal(rng.uniform(5.0, 8.0), 1.0, n)
        module_of[f"N{k:02d}"] = "noise"
    for k in range(1, 3):
        data[f"L{k:02d}"] = rng.normal(6.0, 0.15, n)
        module_of[f"L{k:02d}"] = "low range"
    expr = pd.DataFrame(data).round(3)
    expr.insert(0, "Group", group)
    expr.insert(0, "SampleID", [f"S{k:03d}" for k in range(1, n + 1)])
    truth = pd.DataFrame({"gene": list(module_of), "module": list(module_of.values())})
    return expr, truth


def write_workbook(expr, truth, path):
    wb = Workbook()
    arial, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)

    ws = wb.active
    ws.title = "expression"
    ws.append(list(expr.columns))
    for row in expr.itertuples(index=False):
        ws.append(list(row))
    ws.freeze_panes = "C2"

    readme = wb.create_sheet("README")
    lines = [
        ("Synthetic test data for birgraph (not real measurements, gene names are placeholders).", True),
        ("", False),
        ("Sheet 'expression': one row per sample, one column per gene; values are on a log2-like scale.", False),
        ("Column SampleID holds sample IDs; column Group is an optional annotation (Tumour or Normal).", False),
        ("To use your own data, keep the same layout: sample ID first, optional annotation columns, then genes.", False),
        ("", False),
        ("Planted structure", True),
    ]
    for text, strong in lines:
        readme.append([text])
        readme.cell(readme.max_row, 1).font = bold if strong else arial
    readme.append(["Relation", "Type", "Meaning"])
    for cell in readme[readme.max_row]:
        cell.font = bold
    for rec in PLANTED:
        readme.append(list(rec))
        for cell in readme[readme.max_row]:
            cell.font = arial
    readme.column_dimensions["A"].width = 22
    readme.column_dimensions["B"].width = 16
    readme.column_dimensions["C"].width = 70

    gt = wb.create_sheet("ground_truth")
    gt.append(["gene", "module"])
    for row in truth.itertuples(index=False):
        gt.append(list(row))

    for sheet in (ws, gt):
        for row in sheet.iter_rows():
            for cell in row:
                cell.font = bold if cell.row == 1 else arial
        for k in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(k)].width = 10
    wb.save(path)
